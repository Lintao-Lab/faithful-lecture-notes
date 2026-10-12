"""Convert a native PPTX into link-free Markdown for LLM consumption.

The original deck remains the visual archive. This converter extracts native
text, tables and SmartArt in slide order, removes hyperlinks and bare URLs, and
inlines reviewed image content supplied through a small JSON mapping. It does
not create or reference image artefacts.
"""

from __future__ import annotations

import argparse
import json
import posixpath
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


_MARKDOWN_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\([^)]*\)")
_MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_BARE_URL_RE = re.compile(r"https?://[^\s<>]+", flags=re.IGNORECASE)
_SPACE_BEFORE_PUNCTUATION_RE = re.compile(r"\s+([,.;:!?])")


@dataclass
class Block:
    top: int
    left: int
    kind: str
    markdown: str


def strip_links(value: str) -> str:
    """Keep visible labels while removing link syntax, targets and bare URLs."""

    value = _MARKDOWN_IMAGE_RE.sub(r"\1", value)
    value = _MARKDOWN_LINK_RE.sub(r"\1", value)
    value = _BARE_URL_RE.sub("", value)
    value = _SPACE_BEFORE_PUNCTUATION_RE.sub(r"\1", value)
    return value


def format_run(run: Any) -> str:
    """Render visible text only; a PowerPoint hyperlink target is never emitted."""

    text = strip_links(run.text or "")
    if not text.strip():
        return ""
    leading = text[: len(text) - len(text.lstrip())]
    trailing = text[len(text.rstrip()) :]
    text = text.strip()
    # Superscript-style or formatted source-note markers add no lecture content.
    if re.fullmatch(r"\[\d+\]", text) and (
        run.font.bold or run.font.underline or getattr(run.hyperlink, "address", None)
    ):
        return leading + trailing
    if run.font.bold:
        text = f"**{text}**"
    if run.font.italic:
        text = f"*{text}*"
    # Underlining on a linked run is presentation chrome, not lecture content.
    if run.font.underline and not getattr(run.hyperlink, "address", None):
        text = f"<u>{text}</u>"
    return leading + text + trailing


def paragraph_markdown(
    paragraph: Any,
    drop_paragraphs: set[str] | None = None,
) -> str:
    plain_text = " ".join((paragraph.text or "").split())
    if drop_paragraphs and plain_text in drop_paragraphs:
        return ""
    content = "".join(format_run(run) for run in paragraph.runs)
    if not content:
        content = strip_links(paragraph.text or "")
    content = content.strip()
    if not content:
        return ""

    prefix = ""
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        tags = {child.tag for child in p_pr}
        indent = "  " * max(0, paragraph.level)
        if qn("a:buAutoNum") in tags:
            prefix = f"{indent}1. "
        elif qn("a:buChar") in tags:
            prefix = f"{indent}- "
    return prefix + content


def text_frame_markdown(shape: Any, drop_paragraphs: set[str] | None = None) -> str:
    paragraphs = [
        paragraph_markdown(p, drop_paragraphs) for p in shape.text_frame.paragraphs
    ]
    return "\n".join(p for p in paragraphs if p)


def table_markdown(shape: Any) -> str:
    rows: list[list[str]] = []
    for row in shape.table.rows:
        cells: list[str] = []
        for cell in row.cells:
            cell_text = "<br>".join(
                p
                for p in (paragraph_markdown(p) for p in cell.text_frame.paragraphs)
                if p
            )
            cells.append(strip_links(cell_text).strip().replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    body = rows[1:]
    output = [f"| {' | '.join(header)} |", f"| {' | '.join(['---'] * width)} |"]
    output.extend(f"| {' | '.join(row)} |" for row in body)
    return "\n".join(output)


def extract_smartart_text(pptx_path: Path) -> dict[int, list[str]]:
    """Read SmartArt diagram-data relationships directly from the PPTX package."""

    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    smartart_by_slide: dict[int, list[str]] = {}
    with zipfile.ZipFile(pptx_path) as archive:
        names = set(archive.namelist())
        slide_names = sorted(
            (name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"\d+", Path(name).stem).group()),
        )
        for slide_name in slide_names:
            slide_number = int(re.search(r"\d+", Path(slide_name).stem).group())
            slide_root = etree.fromstring(archive.read(slide_name))
            relationship_ids = slide_root.xpath(
                ".//dgm:relIds/@r:dm", namespaces=namespaces
            )
            if not relationship_ids:
                continue

            rel_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
            if rel_name not in names:
                continue
            rel_root = etree.fromstring(archive.read(rel_name))
            targets = {
                rel.get("Id"): rel.get("Target")
                for rel in rel_root.xpath(".//rel:Relationship", namespaces=namespaces)
            }
            values: list[str] = []
            for relationship_id in relationship_ids:
                target = targets.get(relationship_id)
                if not target:
                    continue
                target_name = posixpath.normpath(
                    posixpath.join(posixpath.dirname(slide_name), target)
                )
                if target_name not in names:
                    continue
                data_root = etree.fromstring(archive.read(target_name))
                for text in data_root.xpath(".//a:t/text()", namespaces=namespaces):
                    normalised = " ".join(strip_links(str(text)).split())
                    if normalised and normalised not in values:
                        values.append(normalised)
            if values:
                smartart_by_slide[slide_number] = values
    return smartart_by_slide


def flatten_shapes(shapes: Any) -> list[Any]:
    flattened: list[Any] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(flatten_shapes(shape.shapes))
        else:
            flattened.append(shape)
    return flattened


def load_image_content(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError(
            "Image-content JSON must be an object keyed by '<slide>.<image>'."
        )
    result: dict[str, str] = {}
    for key, value in raw.items():
        if not isinstance(key, str) or not re.fullmatch(r"\d+\.\d+", key):
            raise ValueError(f"Invalid image-content key: {key!r}")
        if not isinstance(value, str):
            raise ValueError(f"Image content for {key} must be text.")
        cleaned = strip_links(value).strip()
        if cleaned:
            result[key] = cleaned
    return result


def load_drop_paragraphs(path: Path | None) -> set[str]:
    if path is None:
        return set()
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, list) or not all(isinstance(value, str) for value in raw):
        raise ValueError(
            "Drop-paragraph JSON must be a list of exact visible-text strings."
        )
    return {" ".join(value.split()) for value in raw if value.strip()}


def convert_pptx(
    pptx_path: Path,
    markdown_path: Path,
    report_path: Path,
    image_content: dict[str, str] | None = None,
    drop_paragraphs: set[str] | None = None,
    title: str | None = None,
) -> dict[str, Any]:
    image_content = image_content or {}
    drop_paragraphs = drop_paragraphs or set()
    presentation = Presentation(pptx_path)
    smartart_by_slide = extract_smartart_text(pptx_path)
    lines = [f"# {strip_links(title or pptx_path.stem).strip()}", ""]
    total_text_blocks = 0
    total_tables = 0
    total_pictures = 0
    used_image_keys: set[str] = set()
    slide_reports: list[dict[str, Any]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[Block] = []
        picture_number = 0
        shape_count = 0
        text_count = 0
        table_count = 0
        inline_image_count = 0

        for shape in flatten_shapes(slide.shapes):
            shape_count += 1
            top = int(getattr(shape, "top", 0))
            left = int(getattr(shape, "left", 0))
            if getattr(shape, "has_table", False):
                markdown = table_markdown(shape)
                if markdown:
                    blocks.append(Block(top, left, "table", markdown))
                    table_count += 1
                    total_tables += 1
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_number += 1
                total_pictures += 1
                key = f"{slide_number}.{picture_number}"
                content = image_content.get(key)
                if content:
                    blocks.append(
                        Block(
                            top,
                            left,
                            "image-content",
                            f"Slide {slide_number} image {picture_number} content:\n{content}",
                        )
                    )
                    inline_image_count += 1
                    used_image_keys.add(key)
                continue
            if getattr(shape, "has_text_frame", False):
                markdown = text_frame_markdown(shape, drop_paragraphs)
                if markdown:
                    blocks.append(Block(top, left, "text", markdown))
                    text_count += 1
                    total_text_blocks += 1

        if slide_number in smartart_by_slide:
            blocks.append(
                Block(
                    10**12,
                    0,
                    "smartart",
                    "\n".join(
                        f"- {value}" for value in smartart_by_slide[slide_number]
                    ),
                )
            )
            text_count += 1
            total_text_blocks += 1

        blocks.sort(key=lambda block: (block.top, block.left))
        lines.extend([f"## Slide {slide_number}", ""])
        for block in blocks:
            lines.extend([strip_links(block.markdown).strip(), ""])
        slide_reports.append(
            {
                "slide": slide_number,
                "shapes": shape_count,
                "text_blocks": text_count,
                "tables": table_count,
                "pictures": picture_number,
                "inlined_picture_contents": inline_image_count,
            }
        )

    markdown = "\n".join(lines).rstrip() + "\n"
    if _BARE_URL_RE.search(markdown) or "](" in markdown or "![" in markdown:
        raise ValueError(
            "Link-free validation failed; output still contains link syntax."
        )

    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(markdown, encoding="utf-8", newline="\n")
    report = {
        "source": str(pptx_path),
        "output": str(markdown_path),
        "slides": len(presentation.slides),
        "text_blocks": total_text_blocks,
        "tables": total_tables,
        "pictures_seen": total_pictures,
        "picture_contents_inlined": len(used_image_keys),
        "used_image_content_keys": sorted(used_image_keys),
        "unused_image_content_keys": sorted(set(image_content) - used_image_keys),
        "dropped_paragraphs": sorted(drop_paragraphs),
        "link_free": True,
        "slide_details": slide_reports,
    }
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n"
    )
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--markdown", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--image-content-json", type=Path)
    parser.add_argument("--drop-paragraphs-json", type=Path)
    parser.add_argument("--title")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    image_content = load_image_content(args.image_content_json)
    drop_paragraphs = load_drop_paragraphs(args.drop_paragraphs_json)
    report = convert_pptx(
        pptx_path=args.pptx,
        markdown_path=args.markdown,
        report_path=args.report,
        image_content=image_content,
        drop_paragraphs=drop_paragraphs,
        title=args.title,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
