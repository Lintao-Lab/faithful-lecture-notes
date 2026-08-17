from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from lecture_summary_pipeline.pptx_to_llm_markdown import convert_pptx, strip_links


def test_strip_links_keeps_visible_text_and_citations() -> None:
    value = "Read [source label](https://example.test/x) and [2025] REF 714 https://noise.test"
    assert strip_links(value).strip() == "Read source label and [2025] REF 714"


def test_converter_emits_link_free_text_and_inline_picture_content(
    tmp_path: Path,
) -> None:
    source = tmp_path / "deck.pptx"
    markdown = tmp_path / "deck.md"
    report = tmp_path / "report.json"
    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), "white").save(picture)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "Visible source label"
    run.hyperlink.address = "https://example.test/source"
    raw_url = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    raw_url.text = "https://example.test/noise"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(3))
    presentation.save(source)

    result = convert_pptx(
        source,
        markdown,
        report,
        image_content={
            "1.1": "Screenshot heading. More at https://example.test/ignored"
        },
        title="Test deck",
    )
    output = markdown.read_text(encoding="utf-8")

    assert "Visible source label" in output
    assert "Slide 1 image 1 content:" in output
    assert "Screenshot heading." in output
    assert "http" not in output
    assert "](" not in output
    assert "![" not in output
    assert ".png" not in output
    assert result["picture_contents_inlined"] == 1
